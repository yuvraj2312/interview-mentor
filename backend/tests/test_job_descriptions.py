import io
from unittest.mock import patch

from docx import Document

from tests.conftest import signup_and_get_tokens

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

FAKE_ANALYSIS = {
    "required_skills": ["Python", "PostgreSQL"],
    "preferred_skills": ["Docker"],
    "seniority_level": "senior",
    "low_confidence_fields": [],
}


def _create_jd(client, access_token, raw_text="We need a senior Python engineer with PostgreSQL experience."):
    with patch("app.services.job_description_service.analyze_job_description", return_value=FAKE_ANALYSIS):
        return client.post(
            "/job-descriptions",
            json={"raw_text": raw_text},
            headers={"Authorization": f"Bearer {access_token}"},
        )


def test_create_job_description_requires_auth(client):
    response = client.post("/job-descriptions", json={"raw_text": "some jd"})
    assert response.status_code == 401


def test_create_job_description_is_synchronous_and_ready(client):
    tokens = signup_and_get_tokens(client)
    response = _create_jd(client, tokens["access_token"])
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "ready"
    assert body["structured_data"]["required_skills"] == ["Python", "PostgreSQL"]


def test_create_job_description_returns_502_on_malformed_llm_output(client):
    tokens = signup_and_get_tokens(client)
    with patch(
        "app.services.job_description_service.analyze_job_description",
        side_effect=ValueError("malformed"),
    ):
        response = client.post(
            "/job-descriptions",
            json={"raw_text": "some jd"},
            headers={"Authorization": f"Bearer {tokens['access_token']}"},
        )
    assert response.status_code == 502


def test_user_cannot_access_another_users_job_description(client):
    tokens_a = signup_and_get_tokens(client)
    tokens_b = signup_and_get_tokens(client)

    created = _create_jd(client, tokens_a["access_token"]).json()

    response = client.get(
        f"/job-descriptions/{created['id']}",
        headers={"Authorization": f"Bearer {tokens_b['access_token']}"},
    )
    assert response.status_code == 404


def test_patch_job_description_updates_structured_data(client):
    tokens = signup_and_get_tokens(client)
    created = _create_jd(client, tokens["access_token"]).json()

    response = client.patch(
        f"/job-descriptions/{created['id']}",
        json={"structured_data": {"required_skills": ["Python"], "preferred_skills": [], "seniority_level": "mid"}},
        headers={"Authorization": f"Bearer {tokens['access_token']}"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["structured_data"]["seniority_level"] == "mid"
    assert body["low_confidence_fields"] == []


def _sample_jd_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Senior Backend Engineer")
    doc.add_paragraph("We need a senior Python engineer with PostgreSQL experience.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_job_description_requires_auth(client):
    response = client.post(
        "/job-descriptions/upload", files={"file": ("jd.docx", _sample_jd_docx_bytes(), DOCX_CONTENT_TYPE)}
    )
    assert response.status_code == 401


def test_upload_job_description_rejects_unsupported_content_type(client):
    tokens = signup_and_get_tokens(client)
    response = client.post(
        "/job-descriptions/upload",
        files={"file": ("jd.txt", b"not a jd", "text/plain")},
        headers={"Authorization": f"Bearer {tokens['access_token']}"},
    )
    assert response.status_code == 400


def test_upload_job_description_extracts_text_and_analyzes_synchronously(client):
    tokens = signup_and_get_tokens(client)
    with patch("app.services.job_description_service.analyze_job_description", return_value=FAKE_ANALYSIS):
        response = client.post(
            "/job-descriptions/upload",
            files={"file": ("jd.docx", _sample_jd_docx_bytes(), DOCX_CONTENT_TYPE)},
            headers={"Authorization": f"Bearer {tokens['access_token']}"},
        )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "ready"
    assert body["structured_data"]["required_skills"] == ["Python", "PostgreSQL"]
    assert "Senior Backend Engineer" in body["raw_text"]
    assert "PostgreSQL" in body["raw_text"]


def test_list_job_descriptions_empty_for_new_user(client):
    tokens = signup_and_get_tokens(client)
    response = client.get("/job-descriptions", headers={"Authorization": f"Bearer {tokens['access_token']}"})
    assert response.status_code == 200
    assert response.json() == []


def test_list_job_descriptions_requires_auth(client):
    response = client.get("/job-descriptions")
    assert response.status_code == 401


def test_list_job_descriptions_excludes_other_users_job_descriptions(client):
    tokens_a = signup_and_get_tokens(client)
    tokens_b = signup_and_get_tokens(client)
    _create_jd(client, tokens_a["access_token"])

    response = client.get("/job-descriptions", headers={"Authorization": f"Bearer {tokens_b['access_token']}"})
    assert response.status_code == 200
    assert response.json() == []


def test_list_job_descriptions_ordered_newest_first_with_preview(client):
    tokens = signup_and_get_tokens(client)
    access_token = tokens["access_token"]

    long_text = "We need a senior Python engineer. " * 10
    first = _create_jd(client, access_token, raw_text="short jd").json()
    second = _create_jd(client, access_token, raw_text=long_text).json()

    response = client.get("/job-descriptions", headers={"Authorization": f"Bearer {access_token}"})
    assert response.status_code == 200
    body = response.json()
    assert len(body) == 2
    assert body[0]["id"] == second["id"]
    assert body[1]["id"] == first["id"]
    assert body[0]["raw_text_preview"] == long_text[:160] + "…"
    assert body[1]["raw_text_preview"] == "short jd"
    assert "raw_text" not in body[0]


def test_delete_job_description_requires_auth(client):
    tokens = signup_and_get_tokens(client)
    created = _create_jd(client, tokens["access_token"]).json()
    response = client.delete(f"/job-descriptions/{created['id']}")
    assert response.status_code == 401


def test_delete_job_description_hides_it_from_list_but_get_by_id_still_works(client):
    tokens = signup_and_get_tokens(client)
    access_token = tokens["access_token"]
    created = _create_jd(client, access_token).json()

    response = client.delete(
        f"/job-descriptions/{created['id']}", headers={"Authorization": f"Bearer {access_token}"}
    )
    assert response.status_code == 204

    list_response = client.get("/job-descriptions", headers={"Authorization": f"Bearer {access_token}"})
    assert list_response.json() == []

    get_response = client.get(
        f"/job-descriptions/{created['id']}", headers={"Authorization": f"Bearer {access_token}"}
    )
    assert get_response.status_code == 200
    assert get_response.json()["id"] == created["id"]


def test_delete_job_description_is_scoped_to_owner(client):
    tokens_a = signup_and_get_tokens(client)
    tokens_b = signup_and_get_tokens(client)
    created = _create_jd(client, tokens_a["access_token"]).json()

    response = client.delete(
        f"/job-descriptions/{created['id']}",
        headers={"Authorization": f"Bearer {tokens_b['access_token']}"},
    )
    assert response.status_code == 404
