import io

from docx import Document

from tests.conftest import signup_and_get_tokens

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Test Candidate - Backend Engineer")
    doc.add_paragraph("Skills: Python, FastAPI, PostgreSQL")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload_resume(client, access_token):
    return client.post(
        "/resumes",
        files={"file": ("resume.docx", _sample_docx_bytes(), DOCX_CONTENT_TYPE)},
        headers={"Authorization": f"Bearer {access_token}"},
    )


def test_upload_resume_requires_auth(client):
    response = client.post("/resumes", files={"file": ("resume.docx", _sample_docx_bytes(), DOCX_CONTENT_TYPE)})
    assert response.status_code == 401


def test_upload_resume_rejects_unsupported_content_type(client):
    tokens = signup_and_get_tokens(client)
    response = client.post(
        "/resumes",
        files={"file": ("resume.txt", b"not a resume", "text/plain")},
        headers={"Authorization": f"Bearer {tokens['access_token']}"},
    )
    assert response.status_code == 400


def test_upload_resume_creates_uploaded_row(client):
    tokens = signup_and_get_tokens(client)
    response = _upload_resume(client, tokens["access_token"])
    assert response.status_code == 202, response.text
    body = response.json()
    assert body["status"] == "uploaded"
    assert body["original_filename"] == "resume.docx"
    assert body["structured_data"] is None


def test_get_resume_requires_auth(client):
    tokens = signup_and_get_tokens(client)
    uploaded = _upload_resume(client, tokens["access_token"]).json()
    response = client.get(f"/resumes/{uploaded['id']}")
    assert response.status_code == 401


def test_user_cannot_access_another_users_resume(client):
    tokens_a = signup_and_get_tokens(client)
    tokens_b = signup_and_get_tokens(client)

    uploaded = _upload_resume(client, tokens_a["access_token"]).json()

    response = client.get(
        f"/resumes/{uploaded['id']}",
        headers={"Authorization": f"Bearer {tokens_b['access_token']}"},
    )
    assert response.status_code == 404


def test_get_missing_resume_returns_404(client):
    tokens = signup_and_get_tokens(client)
    response = client.get(
        "/resumes/00000000-0000-0000-0000-000000000000",
        headers={"Authorization": f"Bearer {tokens['access_token']}"},
    )
    assert response.status_code == 404


def test_patch_resume_rejects_when_not_ready(client):
    tokens = signup_and_get_tokens(client)
    uploaded = _upload_resume(client, tokens["access_token"]).json()

    response = client.patch(
        f"/resumes/{uploaded['id']}",
        json={"structured_data": {"skills": ["Python"]}},
        headers={"Authorization": f"Bearer {tokens['access_token']}"},
    )
    assert response.status_code == 400


def test_list_resumes_empty_for_new_user(client):
    tokens = signup_and_get_tokens(client)
    response = client.get("/resumes", headers={"Authorization": f"Bearer {tokens['access_token']}"})
    assert response.status_code == 200
    assert response.json() == []


def test_list_resumes_requires_auth(client):
    response = client.get("/resumes")
    assert response.status_code == 401


def test_list_resumes_excludes_other_users_resumes(client):
    tokens_a = signup_and_get_tokens(client)
    tokens_b = signup_and_get_tokens(client)
    _upload_resume(client, tokens_a["access_token"])

    response = client.get("/resumes", headers={"Authorization": f"Bearer {tokens_b['access_token']}"})
    assert response.status_code == 200
    assert response.json() == []


def test_list_resumes_ordered_newest_first(client):
    tokens = signup_and_get_tokens(client)
    access_token = tokens["access_token"]

    first = _upload_resume(client, access_token).json()
    second = _upload_resume(client, access_token).json()

    response = client.get("/resumes", headers={"Authorization": f"Bearer {access_token}"})
    assert response.status_code == 200
    body = response.json()
    assert len(body) == 2
    assert body[0]["id"] == second["id"]
    assert body[1]["id"] == first["id"]
    assert body[0]["original_filename"] == "resume.docx"
    assert body[0]["status"] == "uploaded"
    assert "structured_data" not in body[0]


def test_delete_resume_requires_auth(client):
    tokens = signup_and_get_tokens(client)
    uploaded = _upload_resume(client, tokens["access_token"]).json()
    response = client.delete(f"/resumes/{uploaded['id']}")
    assert response.status_code == 401


def test_delete_resume_hides_it_from_list_but_get_by_id_still_works(client):
    tokens = signup_and_get_tokens(client)
    access_token = tokens["access_token"]
    uploaded = _upload_resume(client, access_token).json()

    response = client.delete(f"/resumes/{uploaded['id']}", headers={"Authorization": f"Bearer {access_token}"})
    assert response.status_code == 204

    list_response = client.get("/resumes", headers={"Authorization": f"Bearer {access_token}"})
    assert list_response.json() == []

    get_response = client.get(f"/resumes/{uploaded['id']}", headers={"Authorization": f"Bearer {access_token}"})
    assert get_response.status_code == 200
    assert get_response.json()["id"] == uploaded["id"]


def test_delete_resume_is_scoped_to_owner(client):
    tokens_a = signup_and_get_tokens(client)
    tokens_b = signup_and_get_tokens(client)
    uploaded = _upload_resume(client, tokens_a["access_token"]).json()

    response = client.delete(
        f"/resumes/{uploaded['id']}",
        headers={"Authorization": f"Bearer {tokens_b['access_token']}"},
    )
    assert response.status_code == 404
