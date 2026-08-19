import io
import uuid
from unittest.mock import patch

from docx import Document

from app.models import Resume, User
from app.services import resume_service
from app.storage import build_resume_object_key, ensure_bucket_exists, upload_resume_file
from tests.conftest import TestSessionLocal, unique_email

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

FAKE_ANALYSIS = {
    "skills": ["Python", "FastAPI"],
    "experience": [],
    "education": [],
    "projects": [],
    "low_confidence_fields": [],
}


def _make_uploaded_resume(db) -> Resume:
    user = User(email=unique_email(), hashed_password="x")
    db.add(user)
    db.flush()

    resume = Resume(
        user_id=user.id,
        original_filename="resume.docx",
        content_type=DOCX_CONTENT_TYPE,
        storage_key="",
    )
    db.add(resume)
    db.flush()

    doc = Document()
    doc.add_paragraph("Sample Candidate")
    doc.add_paragraph("Skills: Python, FastAPI")
    buf = io.BytesIO()
    doc.save(buf)

    ensure_bucket_exists()
    key = build_resume_object_key(user.id, resume.id, "resume.docx")
    upload_resume_file(buf.getvalue(), key, DOCX_CONTENT_TYPE)
    resume.storage_key = key
    db.add(resume)
    db.commit()
    db.refresh(resume)
    return resume


def test_process_uploaded_resume_marks_ready():
    db = TestSessionLocal()
    try:
        resume = _make_uploaded_resume(db)
        with patch("app.services.resume_service.analyze_resume", return_value=FAKE_ANALYSIS):
            resume_service.process_uploaded_resume(db, resume.id)

        db.refresh(resume)
        assert resume.status == "ready"
        assert resume.structured_data["skills"] == ["Python", "FastAPI"]
        assert resume.extracted_text is not None
        assert resume.low_confidence_fields == []
    finally:
        db.close()


def test_process_uploaded_resume_marks_failed_on_analyzer_error():
    db = TestSessionLocal()
    try:
        resume = _make_uploaded_resume(db)
        with patch("app.services.resume_service.analyze_resume", side_effect=ValueError("bad llm output")):
            resume_service.process_uploaded_resume(db, resume.id)

        db.refresh(resume)
        assert resume.status == "failed"
        assert resume.error_message == "bad llm output"
    finally:
        db.close()


def test_process_uploaded_resume_unknown_id_is_noop():
    db = TestSessionLocal()
    try:
        resume_service.process_uploaded_resume(db, uuid.uuid4())
    finally:
        db.close()
